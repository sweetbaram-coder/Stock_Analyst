import openpyxl

wb = openpyxl.load_workbook("output/news_einfomax_co_kr/fx_news_20260904_141600/crawl_result.xlsx")
ws = wb.active

headers = [cell.value for cell in ws[1]]
print("Headers:", headers)
print(f"Total Rows: {ws.max_row} (Data: {ws.max_row - 1})")

print("\n--- First 3 Records ---")
for r in range(2, 5):
    row_vals = [cell.value for cell in ws[r]]
    print(f"Row {r-1}: {row_vals[0]} | {row_vals[1]} | {row_vals[2]} | {row_vals[4]}")

print("\n--- Last 3 Records ---")
for r in range(ws.max_row - 2, ws.max_row + 1):
    row_vals = [cell.value for cell in ws[r]]
    print(f"Row {r-1}: {row_vals[0]} | {row_vals[1]} | {row_vals[2]} | {row_vals[4]}")
