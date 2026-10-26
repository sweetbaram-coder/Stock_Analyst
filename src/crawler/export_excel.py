"""src/crawler/export_excel.py — 수집 및 분석 데이터를 엑셀 파일로 출력
참고자료(web-crawler/scripts/export_excel.py) 패턴 확장:
- 복수 시트 지원 (테마 랭킹, 주도주 후보, 대장주 기술지표)
- 전문적인 스타일링 (헤더 스타일, 열 너비 자동 조정, 오토필터, 숫자 서식)
"""
import os
from typing import Dict, List, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def export_to_excel(data: List[Dict[str, Any]], filepath: str, sheet_name: str = "수집 데이터") -> Optional[str]:
    """단일 시트 엑셀 내보내기 (기본 호환 함수)"""
    if not data:
        return None
    return export_multi_sheet_excel({sheet_name: data}, filepath)

def export_multi_sheet_excel(sheets_data: Dict[str, List[Dict[str, Any]]], filepath: str) -> Optional[str]:
    """다중 시트 엑셀 내보내기

    Args:
        sheets_data: {"시트명": [ {컬럼: 값, ...}, ... ]}
        filepath: 저장 파일 경로
    """
    if not sheets_data:
        return None

    os.makedirs(os.path.dirname(filepath) or ".", exist_ok=True)
    wb = Workbook()
    # 기본 첫 시트 제거 준비
    first_sheet_created = False

    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    data_font = Font(name="맑은 고딕", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet_name, data in sheets_data.items():
        if not data:
            continue

        if not first_sheet_created:
            ws = wb.active
            ws.title = sheet_name
            first_sheet_created = True
        else:
            ws = wb.create_sheet(title=sheet_name)

        headers = list(data[0].keys())

        # 헤더 행 작성
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # 데이터 행 작성
        for row_idx, item in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                val = item.get(header)
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border

                # 정렬 및 서식
                if isinstance(val, (int, float)):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if isinstance(val, int) or (isinstance(val, float) and val.is_integer()):
                        cell.number_format = "#,##0"
                    else:
                        cell.number_format = "#,##0.00"
                elif isinstance(val, str) and (val.endswith("%") or val.startswith("+") or val.startswith("-")):
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # 열 너비 자동 계산
        for col_idx, header in enumerate(headers, 1):
            max_len = len(str(header).encode('euc-kr', errors='ignore'))
            for row in range(2, min(len(data) + 2, 50)):
                v = ws.cell(row=row, column=col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v).encode('euc-kr', errors='ignore')))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(min(max_len + 3, 50), 10)

        # 오토필터 적용
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(data) + 1}"

    wb.save(filepath)
    return filepath
